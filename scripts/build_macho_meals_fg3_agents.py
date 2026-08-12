"""
Build the 5 missing Macho Meals FG3 agents (completes the 22-participant roster).

Deterministic ETL, no API calls. Companion to scripts/build_fg_agents.py, which
skipped FG3 (dash-PID rows). The researcher has since resolved the PID-matching
gap by random 1:1 assignment of the 5 real FG3 survey rows to the 5 known FG3
transcript speaker names; manifest v5 now carries the real names in FG3's
`Pseudonym` column.

Psychometric construct/direction/scale boilerplate is copied VERBATIM from the
on-disk mm_fg1_amir.json rather than retyped, so the fixed construct text stays
byte-identical to the other 17 agents.

Usage: py scripts/build_macho_meals_fg3_agents.py
"""
import io
import json
import sys
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
AGENT_DIR = ROOT / "agents" / "macho_meals"
TEMPLATE = AGENT_DIR / "mm_fg1_amir.json"
MANIFEST_XLSX = ROOT / "data" / "manifests" / "focus_group_dataset_manifest_v5.xlsx"
FG3_STD = ROOT / "data" / "datasets_transcripts" / "standardized" / "macho_meals" / "fg3"

PK = [
    "masculine_norms",
    "masculinity_of_meat",
    "meat_attachment",
    "dairy_attachment",
    "vegetarianism_threat",
]
FN = ["red_meat", "poultry", "fish", "egg", "dairy"]

LINKAGE_NOTE = (
    "FG3 (the 40-49 age-band focus group) had a PID recording error that made it "
    "impossible to match individual survey responses to individual named transcript "
    "speakers. The 5 real, unaltered FG3 survey rows (demographics + psychometric "
    "scores + consumption frequencies) were assigned to the 5 known FG3 transcript "
    "speaker names (Andrew, Daniel, John, Nick, Paul) by the researcher via random "
    "1:1 pairing, recorded in data/manifests/focus_group_dataset_manifest_v5.xlsx, "
    "replacing the dash placeholders documented in "
    "data/datasets_transcripts/standardized/macho_meals/fg3/baseline_metadata.json. "
    "Every field value below is genuine, unaltered FG3 survey data; only the specific "
    "person-to-row correspondence is arbitrary. Analyses that treat an individual FG3 "
    "agent's psychometric profile as predictive of that same named agent's transcript "
    "statements are not valid at the individual level for this group; FG3 group-level/"
    "aggregate findings are unaffected, since all 5 real rows are genuinely FG3 data."
)

# Same note, but the parenthetical cross-reference differs per destination file:
# in _manifest.json, `excluded_historical` is a real sibling key; in
# baseline_metadata.json it is not — there the local record is `fg3_exclusion_note`,
# written directly above. Keeping these as one shared constant would silently
# reintroduce the dangling reference into baseline_metadata.json on any re-run.
RESOLVED_NOTE_MANIFEST = (
    "The 5 FG3 dash-PID participants (see 'excluded_historical' below) were built once "
    "the researcher resolved the PID-matching gap via random 1:1 assignment of survey "
    "rows to known transcript speaker names. See each mm_fg3_*.json's "
    "study_context.identity_metadata_linkage_note for the full caveat."
)

RESOLVED_NOTE_BASELINE = (
    "The 5 FG3 dash-PID participants (see 'fg3_exclusion_note' above) were built once "
    "the researcher resolved the PID-matching gap via random 1:1 assignment of survey "
    "rows to known transcript speaker names. See each mm_fg3_*.json's "
    "study_context.identity_metadata_linkage_note for the full caveat."
)


def write_json(path: Path, data) -> None:
    path.write_text(
        json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )


# ── Lift the fixed construct boilerplate verbatim from the existing template ──

template = json.loads(TEMPLATE.read_text(encoding="utf-8"))
PSYCH = {}
for pk in PK:
    entry = dict(template["psychometric_scores"][pk])
    entry.pop("value", None)
    entry.pop("provenance", None)
    PSYCH[pk] = entry

# ── Read the 5 FG3 rows from the manifest ────────────────────────────────────

wb = openpyxl.load_workbook(MANIFEST_XLSX, data_only=True, read_only=True)
rows = list(wb["DS03_MACHO_MEALS_UK"].iter_rows(values_only=True))

built = []
for r in rows[1:]:
    if r[0] is None or r[1] is None or r[2] is None:
        continue
    try:
        fg = int(r[0])
    except (ValueError, TypeError):
        continue
    if fg != 3:
        continue

    name = str(r[1]).strip()
    aid = f"mm_fg3_{name.lower()}"

    dem = {
        "name": name,
        "age": int(r[2]),
        "gender": str(r[3]).strip(),
        "diet": str(r[6]).strip(),
        "location": {
            "urban_rural": str(r[5]).strip(),
            "region": str(r[4]).strip(),
            "country": "UK",
        },
    }

    fc = {fn: str(r[12 + i]).strip() for i, fn in enumerate(FN)}

    ps = {}
    for i, pk in enumerate(PK):
        entry = dict(PSYCH[pk])
        entry["value"] = float(r[7 + i])
        entry["provenance"] = "observed"
        ps[pk] = entry

    diet_s = dem["diet"]
    notes_text = (
        "Recruited as a regular meat eater."
        if diet_s.lower() == "meat eater"
        else f"Self-identified {diet_s.lower()}."
    )

    intro = str(r[17]).strip() if r[17] else None
    if intro and intro != "None":
        raise SystemExit(
            f"Unexpected non-empty intro cell for FG3 row {name!r} — §1.5 of the "
            "instructions asserts all 5 FG3 intro cells are empty. Stopping."
        )

    agent = {
        "schema_version": "fg_agents_v1",
        "agent_id": aid,
        "language": "en",
        "field_provenance": {
            "language": "observed",
            "persona.demographics.name": "observed",
            "persona.demographics.age": "observed",
            "persona.demographics.gender": "observed",
            "persona.demographics.diet": "observed",
            "persona.demographics.location.urban_rural": "observed",
            "persona.demographics.location.region": "observed",
            "persona.demographics.location.country": "derived",
            "persona.food_consumption": "observed",
            "psychometric_scores": "observed",
            "simulation_config.notes": "derived",
            "study_context.identity_metadata_linkage": "researcher_declared",
        },
        "persona": {"demographics": dem, "food_consumption": fc},
        "psychometric_scores": ps,
        "study_context": {
            "dataset": "DS03_MACHO_MEALS_UK",
            "focus_group": "FG3",
            "identity_metadata_linkage": "researcher_random_assignment",
            "identity_metadata_linkage_note": LINKAGE_NOTE,
        },
        "simulation_config": {
            "model": "claude-haiku-4-5-20251001",
            "max_tokens": 400,
            "notes": notes_text,
        },
    }
    # No "opening_intro" key: intro cell empty (sf_fg3_angel omission pattern).

    write_json(AGENT_DIR / f"{aid}.json", agent)
    built.append(aid)
    print(f"  wrote {aid}.json")

if len(built) != 5:
    raise SystemExit(f"Expected 5 FG3 rows, got {len(built)}: {built}")

# ── _manifest.json ───────────────────────────────────────────────────────────

manifest_path = AGENT_DIR / "_manifest.json"
old = json.loads(manifest_path.read_text(encoding="utf-8"))
# The first run reads the original "excluded" record; every later run reads the
# "excluded_historical" this script already renamed it to. Accept either, so the
# historical wording is carried through verbatim instead of KeyError-ing on re-run.
excluded_historical = dict(old.get("excluded") or old["excluded_historical"])
excluded_historical["status"] = "superseded — see fg3_resolution above"

write_json(
    manifest_path,
    {
        "dataset": "DS03_MACHO_MEALS_UK",
        "schema_version": "fg_agents_v1",
        "agents_built": 22,
        "intro_eligible": 17,
        "fg3_resolution": {
            "resolved": True,
            "resolved_note": RESOLVED_NOTE_MANIFEST,
            "agent_ids": sorted(built),
        },
        "excluded_historical": excluded_historical,
    },
)
print("  updated _manifest.json (agents_built: 17 -> 22)")

# ── FG3 standardized-transcript metadata ─────────────────────────────────────

name_to_aid = {aid.split("_")[-1].capitalize(): aid for aid in built}

recon_path = FG3_STD / "identity_reconciliation.json"
recon = json.loads(recon_path.read_text(encoding="utf-8"))
for e in recon:
    e["agent_id"] = name_to_aid[e["transcript_speaker"]]
    e["matched"] = True
write_json(recon_path, recon)
print("  updated fg3/identity_reconciliation.json")

pm_path = FG3_STD / "participant_metadata.json"
pm = json.loads(pm_path.read_text(encoding="utf-8"))
for e in pm:
    if e.get("speaker_role") == "participant":
        e["agent_id"] = name_to_aid[e["speaker_name"]]
write_json(pm_path, pm)
print("  updated fg3/participant_metadata.json")

bm_path = FG3_STD / "baseline_metadata.json"
bm = json.loads(bm_path.read_text(encoding="utf-8"))
bm["agents_matched"] = True
bm["agents_matched_count"] = 5
bm["no_matched_agents"] = False
bm["fg3_resolution_note"] = RESOLVED_NOTE_BASELINE  # fg3_exclusion_note preserved as-is
write_json(bm_path, bm)
print("  updated fg3/baseline_metadata.json")

print(f"\nDONE — {len(built)} FG3 agents built: {', '.join(sorted(built))}")
